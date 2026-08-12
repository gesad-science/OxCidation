import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from review_workbook import generate_review_workbook


class ReviewWorkbookTests(unittest.TestCase):
    def test_generates_analysis_sheets_and_relative_links(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            experiment = Path(temp_dir)
            (experiment / "manifest.json").write_text(
                json.dumps(
                    {
                        "experiment_id": "experiment",
                        "model": {"provider": "openai", "id": "model"},
                        "sampling": {"sample_size": 1},
                        "prompts": [{"prompt_id": "direct"}],
                    }
                ),
                encoding="utf-8",
            )
            result = {
                "review_id": "blind-id",
                "snippet_id": "s1",
                "problem_id": "p1",
                "prompt_id": "direct",
                "visible_suite_status": "ready",
                "visible_suite_c_compile_profile": "gnu11",
                "prompt_sha256": "hash",
                "initial_compile_status": "success",
                "initial_visible_test_status": "success",
                "initial_judge_status": "ACCEPTED",
                "final_status": "success",
                "judge_status": "ACCEPTED",
                "repair_attempts": "0",
                "result_path": "model_outputs/s1/direct/result.json",
                "initial_code_path": "model_outputs/s1/direct/attempts/attempt-00-initial.rs",
                "final_code_path": "model_outputs/s1/direct/translated.rs",
                "agent_interactions_path": "model_outputs/s1/direct/agent_interactions.json",
            }
            skipped_result = {
                **result,
                "review_id": "skipped-id",
                "snippet_id": "s2",
                "judge_status": "SKIPPED",
            }
            not_reached_result = {
                **result,
                "review_id": "not-reached-id",
                "snippet_id": "s3",
                "judge_status": "not_reached",
            }
            failed_result = {
                **result,
                "review_id": "failed-id",
                "snippet_id": "s4",
                "judge_status": "WRONG_ANSWER",
            }
            self._write_csv(
                experiment / "results.csv",
                [result, skipped_result, not_reached_result, failed_result],
            )
            self._write_csv(
                experiment / "attempts.csv",
                [
                    {
                        "review_id": "blind-id",
                        "attempt_number": "0",
                    }
                ],
            )

            path = generate_review_workbook(experiment)
            workbook = load_workbook(path)

            self.assertNotIn("Prompt Mapping", workbook.sheetnames)
            self.assertEqual(workbook["Review Queue"]["A2"].value, "blind-id")
            self.assertNotIn("Reviews", workbook.sheetnames)
            self.assertNotIn("Adjudication", workbook.sheetnames)
            self.assertIn(
                "Prompt ID",
                [cell.value for cell in workbook["Review Queue"][1]],
            )
            self.assertEqual(workbook["Automatic Results"].max_column, 24)
            self.assertEqual(workbook["Attempts"].max_column, 19)
            queue_headers = [
                cell.value for cell in workbook["Review Queue"][1]
            ]
            self.assertIn("Visible Suite Status", queue_headers)
            self.assertIn("Visible Suite C Compile Profile", queue_headers)
            overview_headers = [
                cell.value for cell in workbook["Overview"][8]
            ]
            self.assertIn("Judge failures", overview_headers)
            self.assertIn("Judge skipped", overview_headers)
            self.assertIn("Judge not reached", overview_headers)
            overview_values = {
                header: workbook["Overview"].cell(9, index + 1).value
                for index, header in enumerate(overview_headers)
            }
            self.assertEqual(overview_values["Runs"], 4)
            self.assertEqual(overview_values["Judge failures"], 1)
            self.assertEqual(overview_values["Judge skipped"], 1)
            self.assertEqual(overview_values["Judge not reached"], 1)
            self.assertIsNotNone(workbook["Artifact Index"]["D2"].hyperlink)
            self.assertFalse(
                any(
                    cell.hyperlink
                    for row in workbook["Review Queue"].iter_rows()
                    for cell in row
                )
            )
            hyperlinks = [
                cell.hyperlink.target
                for sheet in workbook.worksheets
                for row in sheet.iter_rows()
                for cell in row
                if cell.hyperlink
            ]
            self.assertTrue(hyperlinks)
            self.assertTrue(
                all(
                    not link.startswith("#")
                    and not Path(link).is_absolute()
                    for link in hyperlinks
                )
            )

    @staticmethod
    def _write_csv(path: Path, rows: list[dict]) -> None:
        with path.open("w", newline="", encoding="utf-8") as output:
            writer = csv.DictWriter(output, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
