"""Create an Excel workbook for benchmark analysis and artifact navigation."""

import argparse
import json
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from benchmark_reporting import read_csv


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LINK_FONT = Font(color="0563C1", underline="single")

AUTOMATIC_RESULT_FIELDS = [
    "review_id",
    "snippet_id",
    "problem_id",
    "prompt_id",
    "initial_compile_status",
    "initial_visible_test_status",
    "initial_judge_status",
    "final_status",
    "failure_category",
    "visible_test_status",
    "visible_suite_status",
    "visible_suite_review_status",
    "visible_suite_c_compile_profile",
    "visible_suite_review_count",
    "visible_suite_generation_attempts",
    "visible_suite_deterministic_rejections",
    "visible_suite_validator_replacements",
    "visible_suite_unresolved_replacements",
    "visible_cases_total",
    "visible_cases_passed",
    "judge_status",
    "judge_cases_passed",
    "judge_cases_failed",
    "judge_first_failure_verdict",
    "repair_attempts",
    "duration_seconds",
]

ATTEMPT_REVIEW_FIELDS = [
    "review_id",
    "snippet_id",
    "problem_id",
    "prompt_id",
    "attempt_number",
    "attempt_kind",
    "compile_status",
    "compiler_output_excerpt",
    "visible_test_status",
    "visible_cases_total",
    "visible_cases_passed",
    "visible_cases_failed",
    "visible_verdict_counts",
    "validator_test_assessment",
    "validator_diagnosis",
    "validator_repair_guidance",
    "repair_failure_category",
    "initial_judge_status",
    "final_judge_status",
]

ARTIFACT_FIELDS = [
    ("source", lambda row, directory: directory / "source.c"),
    ("raw_model_output", lambda row, directory: Path(row.get("raw_output_path", ""))),
    (
        "initial_translation",
        lambda row, directory: Path(row.get("initial_code_path", "")),
    ),
    ("final_translation", lambda row, directory: Path(row.get("final_code_path", ""))),
    (
        "visible_test_report",
        lambda row, directory: Path(row.get("test_report_path", "")),
    ),
    ("visible_test_suite", lambda row, directory: directory / "test_suite.json"),
    (
        "visible_suite_preparation",
        lambda row, directory: Path(
            row.get("visible_suite_preparation_history_path", "")
        ),
    ),
    (
        "validator_report",
        lambda row, directory: Path(row.get("validator_report_path", "")),
    ),
    (
        "agent_interactions",
        lambda row, directory: Path(row.get("agent_interactions_path", "")),
    ),
    ("attempt_summary", lambda row, directory: Path(row.get("attempts_path", ""))),
    (
        "initial_judge_report",
        lambda row, directory: directory / "initial_judge_comparison.json",
    ),
    ("final_judge_report", lambda row, directory: directory / "judge_comparison.json"),
    ("complete_result", lambda row, directory: Path(row.get("result_path", ""))),
]


def generate_review_workbook(
    experiment_dir: Path,
    output_path: Path | None = None,
) -> Path:
    manifest = json.loads(
        (experiment_dir / "manifest.json").read_text(encoding="utf-8")
    )
    results = read_csv(experiment_dir / "results.csv")
    attempts = read_csv(experiment_dir / "attempts.csv")
    output_path = output_path or experiment_dir / "review_workbook.xlsx"

    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_overview(workbook, manifest, results)
    _write_review_queue(workbook, experiment_dir, results)
    _write_artifact_index(workbook, results)
    _write_table(
        workbook,
        "Automatic Results",
        AUTOMATIC_RESULT_FIELDS,
        results,
    )
    _write_table(
        workbook,
        "Attempts",
        ATTEMPT_REVIEW_FIELDS,
        attempts,
    )
    workbook.save(output_path)
    return output_path


def _write_overview(workbook: Workbook, manifest: dict, results: list[dict]) -> None:
    sheet = workbook.create_sheet("Overview")
    sheet.append(["Experiment", manifest.get("experiment_id", "")])
    sheet.append(["Model provider", manifest.get("model", {}).get("provider", "")])
    sheet.append(["Model", manifest.get("model", {}).get("id", "")])
    sheet.append(["Programs", manifest.get("sampling", {}).get("sample_size", 0)])
    sheet.append(["Prompts", len(manifest.get("prompts", []))])
    sheet.append(["Completed combinations", len(results)])
    sheet.append([])
    sheet.append(
        [
            "Prompt",
            "Runs",
            "Initial accepted",
            "Final accepted",
            "Repaired",
            "Judge failures",
            "Judge skipped",
            "Judge not reached",
        ]
    )
    prompt_ids = sorted({row["prompt_id"] for row in results})
    for prompt_id in prompt_ids:
        prompt_rows = [row for row in results if row["prompt_id"] == prompt_id]
        sheet.append(
            [
                prompt_id,
                len(prompt_rows),
                _count(prompt_rows, "initial_judge_status", "ACCEPTED"),
                _count(prompt_rows, "judge_status", "ACCEPTED"),
                sum(int(row.get("repair_attempts", 0)) > 0 for row in prompt_rows),
                sum(_is_judge_failure(row.get("judge_status")) for row in prompt_rows),
                sum(_normalized_judge_status(row.get("judge_status")) == "SKIPPED" for row in prompt_rows),
                sum(_judge_not_reached(row.get("judge_status")) for row in prompt_rows),
            ]
        )
    _style_sheet(sheet, freeze="A8", filter_row=8)
    sheet.column_dimensions["A"].width = 28
    for column in ("B", "C", "D", "E", "F", "G", "H"):
        sheet.column_dimensions[column].width = 20


def _count(rows: list[dict], field: str, value: str) -> int:
    return Counter(row.get(field) for row in rows)[value]


def _normalized_judge_status(status: object) -> str:
    return str(status or "").strip().upper()


def _judge_not_reached(status: object) -> bool:
    return _normalized_judge_status(status) in {"", "NOT_REACHED"}


def _is_judge_failure(status: object) -> bool:
    normalized = _normalized_judge_status(status)
    return normalized not in {"", "ACCEPTED", "SKIPPED", "NOT_REACHED"}


def _write_review_queue(
    workbook: Workbook,
    experiment_dir: Path,
    results: list[dict],
) -> None:
    fields = [
        "review_id",
        "snippet_id",
        "problem_id",
        "prompt_id",
        "visible_suite_status",
        "visible_suite_c_compile_profile",
        "initial_compile_status",
        "initial_visible_test_status",
        "initial_judge_status",
        "final_status",
        "final_judge_status",
        "repair_attempts",
        "failure_category",
        "visible_failure_example",
        "judge_failure_example",
    ]
    sheet = workbook.create_sheet("Review Queue", 1)
    sheet.append([_header_label(field) for field in fields])
    for row in sorted(results, key=lambda item: item["review_id"]):
        visible_evidence, judge_evidence = _failure_evidence(
            experiment_dir, row
        )
        sheet.append(
            [
                row["review_id"],
                row["snippet_id"],
                row["problem_id"],
                row["prompt_id"],
                row.get("visible_suite_status", ""),
                row.get("visible_suite_c_compile_profile", ""),
                row["initial_compile_status"],
                row["initial_visible_test_status"],
                row["initial_judge_status"],
                row["final_status"],
                row["judge_status"],
                row["repair_attempts"],
                row.get("failure_category", ""),
                visible_evidence,
                judge_evidence,
            ]
        )
    _style_sheet(sheet, freeze="A2", filter_row=1)
    _set_widths(
        sheet,
        {
            "A": 20,
            "B": 18,
            "C": 14,
            "D": 28,
            "E": 22,
            "F": 22,
            "G": 20,
            "H": 24,
            "I": 20,
            "J": 16,
            "K": 20,
            "L": 16,
            "M": 18,
            "N": 42,
            "O": 42,
        },
        default=18,
    )


def _write_artifact_index(workbook: Workbook, results: list[dict]) -> None:
    fields = [
        "review_id",
        "snippet_id",
        "problem_id",
        *(name for name, _ in ARTIFACT_FIELDS),
    ]
    sheet = workbook.create_sheet("Artifact Index")
    sheet.append([_header_label(field) for field in fields])
    for row in sorted(results, key=lambda item: item["review_id"]):
        sheet.append(
            [
                row["review_id"],
                row["snippet_id"],
                row["problem_id"],
                *(["Open"] * len(ARTIFACT_FIELDS)),
            ]
        )
        artifact_dir = Path(row["result_path"]).parent
        for offset, (_, path_builder) in enumerate(ARTIFACT_FIELDS, start=4):
            _set_file_link(
                sheet.cell(sheet.max_row, offset),
                path_builder(row, artifact_dir),
            )
    _style_sheet(sheet, freeze="A2", filter_row=1)
    _set_widths(
        sheet,
        {"A": 20, "B": 18, "C": 14},
        default=18,
    )


def _failure_evidence(
    experiment_dir: Path,
    row: dict,
) -> tuple[str, str]:
    test_report = _read_json(
        experiment_dir / row.get("test_report_path", "")
    )
    visible = test_report.get("failure_examples", [])
    visible_evidence = _compact_json(visible[0]) if visible else ""

    artifact_dir = Path(row.get("result_path", "")).parent
    judge_report = _read_json(
        experiment_dir / artifact_dir / "judge_comparison.json"
    )
    judge_failure = (
        judge_report.get("rust", {}).get("judge", {}).get("first_failure", {})
    )
    return visible_evidence, _compact_json(judge_failure)


def _read_json(path: Path) -> dict:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}


def _compact_json(value, limit: int = 1200) -> str:
    if not value:
        return ""
    text = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return text if len(text) <= limit else f"{text[:limit - 3]}..."


def _write_table(
    workbook: Workbook,
    name: str,
    fields: list[str],
    rows: list[dict],
) -> None:
    sheet = workbook.create_sheet(name)
    if fields:
        sheet.append([_header_label(field) for field in fields])
    for row in rows:
        sheet.append([_excel_value(row.get(field, "")) for field in fields])
    _style_sheet(sheet, freeze="A2", filter_row=1)
    _set_widths(sheet, {}, default=20)


def _excel_value(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    if isinstance(value, str) and len(value) > 32767:
        return f"{value[:32764]}..."
    return value


def _set_file_link(cell, target: Path) -> None:
    if target == Path(".") or target.is_absolute():
        cell.value = "Unavailable"
        return
    cell.hyperlink = target.as_posix()
    cell.font = LINK_FONT


def _header_label(field: str) -> str:
    return field.replace("_", " ").title().replace(" Id", " ID")


def _style_sheet(sheet, freeze: str, filter_row: int) -> None:
    sheet.freeze_panes = freeze
    if sheet.max_column:
        sheet.auto_filter.ref = (
            f"A{filter_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
        )
    for cell in sheet[filter_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in sheet.iter_rows(min_row=filter_row + 1):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _set_widths(sheet, widths: dict[str, int], default: int) -> None:
    for index in range(1, sheet.max_column + 1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = widths.get(letter, default)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Generate an Excel workbook for manual benchmark review."
    )
    parser.add_argument("experiment_dir", type=Path)
    parser.add_argument("--output", type=Path)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    output = generate_review_workbook(args.experiment_dir, args.output)
    print(f"Review workbook: {output}")


if __name__ == "__main__":
    main()
